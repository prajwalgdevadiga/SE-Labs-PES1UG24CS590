# -*- coding: utf-8 -*-
"""
Builds every text deliverable from content.py:

  docs/01_Requirements_Table.docx / .xlsx / .pdf
  docs/03_UseCase_Flow_UC-01.docx / .pdf

    python build_docs.py
"""
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, KeepTogether, PageTemplate,
                                Paragraph, Spacer, Table, TableStyle)

from content import STUDENT, SCENARIO, REQUIREMENTS, ACTORS, USE_CASES, RELATIONSHIPS, FLOW

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.join(os.path.dirname(HERE), "docs")
os.makedirs(DOCS, exist_ok=True)

HDR_BG = "2F5C9E"
HDR_RGB = colors.HexColor("#2F5C9E")
ALT_RGB = colors.HexColor("#F4F7FC")

COLS = ["Req ID", "Type", "Description (“The system shall …”)", "Priority",
        "Acceptance Criteria (measurable pass/fail)", "Rationale",
        "Comments (peer critique → revision)"]

ROWS = [[r["id"], r["type"], r["description"], r["priority"],
         r["acceptance"], r["rationale"], r["comments"]] for r in REQUIREMENTS]


# ============================================================== docx helpers
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def set_cell(cell, text, bold=False, size=8, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color


def doc_header(doc, title, subtitle=None):
    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(2)
    r = t.add_run(title)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    if subtitle:
        s = doc.add_paragraph()
        s.paragraph_format.space_after = Pt(8)
        rs = s.add_run(subtitle)
        rs.font.size = Pt(9.5)
        rs.italic = True
        rs.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    meta = doc.add_table(rows=2, cols=4)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    pairs = [("Name", STUDENT["name"]), ("SRN", STUDENT["srn"]),
             ("Course", STUDENT["course"]), ("Problem Statement", STUDENT["ps"])]
    for i, (k, v) in enumerate(pairs):
        set_cell(meta.rows[0].cells[i], k, bold=True, size=8,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(meta.rows[0].cells[i], HDR_BG)
        set_cell(meta.rows[1].cells[i], v, size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ============================================================== 1. docx table
def build_requirements_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(0.45))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)

    doc_header(doc, "Deliverable 1 — Requirements Table",
               "%s  •  %s" % (STUDENT["system"], STUDENT["topic"]))

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Scenario: ")
    r.bold = True
    r.font.size = Pt(8.5)
    r2 = p.add_run(SCENARIO)
    r2.font.size = Pt(8.5)

    widths = [Inches(0.62), Inches(0.95), Inches(2.55), Inches(0.52),
              Inches(2.65), Inches(1.75), Inches(1.75)]

    tbl = doc.add_table(rows=1, cols=len(COLS))
    tbl.style = "Table Grid"
    tbl.autofit = False
    for i, c in enumerate(COLS):
        set_cell(tbl.rows[0].cells[i], c, bold=True, size=8,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(tbl.rows[0].cells[i], HDR_BG)

    for n, row in enumerate(ROWS):
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            bold = (i == 0)
            set_cell(cells[i], val, bold=bold, size=7.5)
            if row[0].startswith("NFR"):
                shade(cells[i], "FFF7E8")
            elif n % 2 == 1:
                shade(cells[i], "F4F7FC")

    for row in tbl.rows:
        for i, c in enumerate(row.cells):
            c.width = widths[i]

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(6)
    rn = note.add_run(
        "FR-001 and NFR-001 are the requirements supplied with Problem Statement #35; "
        "FR-002–FR-005 and NFR-002 are student-authored. The Comments column records the "
        "peer-critique round (Lab step 3) and the revision each comment produced."
    )
    rn.font.size = Pt(7.5)
    rn.italic = True
    rn.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # ---- actors & use cases (Lab step 4) on a second page
    doc.add_page_break()
    h = doc.add_paragraph()
    hr = h.add_run("Actors and Use Cases extracted from the approved requirements")
    hr.bold = True
    hr.font.size = Pt(12.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    at = doc.add_table(rows=1, cols=3)
    at.style = "Table Grid"
    for i, c in enumerate(["Actor", "Classification", "Goal in the system"]):
        set_cell(at.rows[0].cells[i], c, bold=True, size=8,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(at.rows[0].cells[i], HDR_BG)
    for a in ACTORS:
        cells = at.add_row().cells
        for i, v in enumerate(a):
            set_cell(cells[i], v, size=8, bold=(i == 0))
    for row in at.rows:
        row.cells[0].width = Inches(1.6)
        row.cells[1].width = Inches(2.0)
        row.cells[2].width = Inches(6.2)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    ut = doc.add_table(rows=1, cols=4)
    ut.style = "Table Grid"
    for i, c in enumerate(["UC ID", "Use Case", "Actor / relationship", "Traces to"]):
        set_cell(ut.rows[0].cells[i], c, bold=True, size=8,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(ut.rows[0].cells[i], HDR_BG)
    for u in USE_CASES:
        cells = ut.add_row().cells
        for i, v in enumerate(u):
            set_cell(cells[i], v, size=8, bold=(i == 0))
    for row in ut.rows:
        row.cells[0].width = Inches(0.8)
        row.cells[1].width = Inches(3.0)
        row.cells[2].width = Inches(3.0)
        row.cells[3].width = Inches(3.0)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    rt = doc.add_table(rows=1, cols=4)
    rt.style = "Table Grid"
    for i, c in enumerate(["Stereotype", "From", "To", "Why"]):
        set_cell(rt.rows[0].cells[i], c, bold=True, size=8,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(rt.rows[0].cells[i], HDR_BG)
    for kind, src, dst, why in RELATIONSHIPS:
        cells = rt.add_row().cells
        set_cell(cells[0], "«%s»" % kind, size=8, bold=True)
        set_cell(cells[1], src, size=8)
        set_cell(cells[2], dst, size=8)
        set_cell(cells[3], why, size=8)

    out = os.path.join(DOCS, "01_Requirements_Table.docx")
    doc.save(out)
    print("wrote docs/01_Requirements_Table.docx")


# ============================================================== 2. xlsx table
def build_requirements_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"

    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "Deliverable 1 - Requirements Table | %s" % STUDENT["system"]
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "%s | %s | %s" % (STUDENT["name"], STUDENT["srn"], STUDENT["ps"])
    ws["A2"].font = Font(size=9, italic=True, color="555555")
    ws["A3"] = "Scenario: %s" % SCENARIO
    ws["A3"].font = Font(size=9, italic=True, color="555555")

    hdr_row = 5
    for i, c in enumerate(COLS, start=1):
        cell = ws.cell(row=hdr_row, column=i, value=c)
        cell.font = Font(bold=True, size=9, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HDR_BG)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border

    for n, row in enumerate(ROWS):
        r = hdr_row + 1 + n
        fill = "FFF7E8" if row[0].startswith("NFR") else ("F4F7FC" if n % 2 else "FFFFFF")
        for i, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.font = Font(size=9, bold=(i == 1))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = border
        ws.row_dimensions[r].height = 78

    for i, w in enumerate([9, 20, 46, 9, 50, 34, 34], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"

    ws2 = wb.create_sheet("Actors & Use Cases")
    ws2.append(["Actor", "Classification", "Goal in the system"])
    for a in ACTORS:
        ws2.append(list(a))
    ws2.append([])
    ws2.append(["UC ID", "Use Case", "Actor / relationship", "Traces to"])
    for u in USE_CASES:
        ws2.append(list(u))
    ws2.append([])
    ws2.append(["Stereotype", "From", "To", "Why"])
    for kind, src, dst, why in RELATIONSHIPS:
        ws2.append(["<<%s>>" % kind, src, dst, why])
    for i, w in enumerate([16, 34, 34, 60], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for r in (1, len(ACTORS) + 3, len(ACTORS) + len(USE_CASES) + 5):
        for c in ws2[r]:
            c.font = Font(bold=True, color="FFFFFF", size=9)
            c.fill = PatternFill("solid", fgColor=HDR_BG)

    out = os.path.join(DOCS, "01_Requirements_Table.xlsx")
    wb.save(out)
    print("wrote docs/01_Requirements_Table.xlsx")


# ============================================================== pdf helpers
def pstyle(name, size, leading, **kw):
    return ParagraphStyle(name, fontName=kw.pop("font", "Helvetica"), fontSize=size,
                          leading=leading, alignment=kw.pop("align", TA_LEFT), **kw)


def make_doc(path, pagesize, title):
    doc = BaseDocTemplate(path, pagesize=pagesize,
                          leftMargin=12 * mm, rightMargin=12 * mm,
                          topMargin=12 * mm, bottomMargin=12 * mm,
                          title=title, author=STUDENT["name"])
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")

    def footer(canv, d):
        canv.saveState()
        canv.setFont("Helvetica", 7)
        canv.setFillColor(colors.HexColor("#777777"))
        canv.drawString(doc.leftMargin, 7 * mm,
                        "%s | %s | %s" % (STUDENT["srn"], STUDENT["name"], STUDENT["course"]))
        canv.drawRightString(pagesize[0] - doc.rightMargin, 7 * mm, "Page %d" % canv.getPageNumber())
        canv.restoreState()

    doc.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=footer)])
    return doc


def title_block(story, title, subtitle):
    story.append(Paragraph(title, pstyle("t", 15, 18, font="Helvetica-Bold",
                                         textColor=colors.HexColor("#1A1A1A"))))
    story.append(Paragraph(subtitle, pstyle("s", 9, 12, font="Helvetica-Oblique",
                                            textColor=colors.HexColor("#555555"))))
    story.append(Spacer(1, 5))
    meta = Table([["Name", "SRN", "Course", "Problem Statement"],
                  [STUDENT["name"], STUDENT["srn"], STUDENT["course"], STUDENT["ps"]]],
                 colWidths=[45 * mm, 32 * mm, 45 * mm, 66 * mm])
    meta.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), HDR_RGB),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(meta)
    story.append(Spacer(1, 8))


# ============================================================== 3. pdf table
def build_requirements_pdf():
    page = landscape(A4)
    doc = make_doc(os.path.join(DOCS, "01_Requirements_Table.pdf"), page,
                   "Requirements Table - %s" % STUDENT["system"])
    body = pstyle("b", 6.9, 8.2)
    bold = pstyle("bb", 6.9, 8.2, font="Helvetica-Bold")
    head = pstyle("h", 7.2, 8.8, font="Helvetica-Bold", textColor=colors.white)

    story = []
    title_block(story, "Deliverable 1 &mdash; Requirements Table",
                "%s &bull; %s" % (STUDENT["system"], STUDENT["topic"]))
    story.append(Paragraph("<b>Scenario:</b> %s" % SCENARIO, pstyle("sc", 7.5, 9.5)))
    story.append(Spacer(1, 6))

    data = [[Paragraph(c, head) for c in COLS]]
    for row in ROWS:
        data.append([Paragraph(row[0], bold)] + [Paragraph(v, body) for v in row[1:]])

    widths = [14 * mm, 24 * mm, 60 * mm, 13 * mm, 63 * mm, 41 * mm, 38 * mm]
    tbl = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), HDR_RGB),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3.5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3.5),
        ("TOPPADDING", (0, 0), (-1, -1), 2.4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4),
    ]
    for i, row in enumerate(ROWS, start=1):
        if row[0].startswith("NFR"):
            style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFF7E8")))
        elif i % 2 == 0:
            style.append(("BACKGROUND", (0, i), (-1, i), ALT_RGB))
    tbl.setStyle(TableStyle(style))
    story.append(tbl)

    story.append(Spacer(1, 4))
    story.append(Paragraph(
        "FR-001 and NFR-001 are the requirements supplied with Problem Statement #35; "
        "FR-002&ndash;FR-005 and NFR-002 are student-authored. The Comments column records the "
        "peer-critique round (Lab step 3) and the revision each comment produced.",
        pstyle("n", 7, 9, font="Helvetica-Oblique", textColor=colors.HexColor("#666666"))))

    # ---- actors & use cases
    from reportlab.platypus import PageBreak
    story.append(PageBreak())
    story.append(Paragraph("Actors and Use Cases extracted from the approved requirements",
                           pstyle("h2", 12.5, 15, font="Helvetica-Bold")))
    story.append(Spacer(1, 6))

    def grid(cols, rows, widths):
        d = [[Paragraph(c, head) for c in cols]]
        for r in rows:
            d.append([Paragraph(str(v), body if i else bold) for i, v in enumerate(r)])
        t = Table(d, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), HDR_RGB),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ALT_RGB]),
        ]))
        return t

    story.append(grid(["Actor", "Classification", "Goal in the system"], ACTORS,
                      [40 * mm, 48 * mm, 165 * mm]))
    story.append(Spacer(1, 10))
    story.append(grid(["UC ID", "Use Case", "Actor / relationship", "Traces to"], USE_CASES,
                      [18 * mm, 70 * mm, 70 * mm, 95 * mm]))
    story.append(Spacer(1, 10))
    story.append(grid(["Stereotype", "From", "To", "Why"],
                      [("«%s»" % k, s, d, w) for k, s, d, w in RELATIONSHIPS],
                      [26 * mm, 22 * mm, 22 * mm, 183 * mm]))

    doc.build(story)
    print("wrote docs/01_Requirements_Table.pdf")


# ============================================================== 4. flow docx
def build_flow_docx():
    doc = Document()
    sec = doc.sections[0]
    for m in ("top_margin", "bottom_margin"):
        setattr(sec, m, Inches(0.5))
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Inches(0.6))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(8.5)

    doc_header(doc, "Deliverable 3 — Use-Case Flow Specification",
               "%s %s  •  %s" % (FLOW["id"], FLOW["name"], STUDENT["system"]))

    mt = doc.add_table(rows=0, cols=2)
    mt.style = "Table Grid"
    for k, v in FLOW["meta"]:
        cells = mt.add_row().cells
        set_cell(cells[0], k, bold=True, size=8)
        shade(cells[0], "F4F7FC")
        set_cell(cells[1], v, size=8)
        cells[0].width = Inches(1.5)
        cells[1].width = Inches(5.9)

    def section(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x2F, 0x5C, 0x9E)

    def numbered(items, size=8.3, indent=0.25):
        for i, it in enumerate(items, start=1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(indent)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(0)
            r = p.add_run("%d.  " % i)
            r.bold = True
            r.font.size = Pt(size)
            r2 = p.add_run(it)
            r2.font.size = Pt(size)

    def plain(items, size=8.1, indent=0.4):
        for it in items:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(indent)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(0)
            r = p.add_run(it)
            r.font.size = Pt(size)

    section("Preconditions")
    numbered(FLOW["preconditions"])
    section("Postconditions")
    numbered(FLOW["postconditions"])
    section("Main Success Scenario")
    numbered(FLOW["main"])
    section("Alternate Flows")
    for alt in FLOW["alternates"]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(alt["title"])
        r.bold = True
        r.font.size = Pt(8.3)
        plain(alt["steps"])
    section("Exception Flow")
    plain([FLOW["exception"]], indent=0.25)

    out = os.path.join(DOCS, "03_UseCase_Flow_UC-01.docx")
    doc.save(out)
    print("wrote docs/03_UseCase_Flow_UC-01.docx")


# ============================================================== 5. flow pdf
def build_flow_pdf():
    doc = make_doc(os.path.join(DOCS, "03_UseCase_Flow_UC-01.pdf"), A4,
                   "Use-Case Flow %s %s" % (FLOW["id"], FLOW["name"]))
    body = pstyle("b", 7.9, 9.9)
    small = pstyle("sm", 7.6, 9.5)
    head = pstyle("h", 7.4, 9, font="Helvetica-Bold", textColor=colors.white)
    sec = pstyle("sec", 9.5, 12, font="Helvetica-Bold",
                 textColor=colors.HexColor("#2F5C9E"), spaceBefore=6, spaceAfter=2)

    story = []
    story.append(Paragraph("Deliverable 3 &mdash; Use-Case Flow Specification",
                           pstyle("t", 14, 17, font="Helvetica-Bold")))
    story.append(Paragraph("%s %s &bull; %s" % (FLOW["id"], FLOW["name"], STUDENT["system"]),
                           pstyle("s", 9, 12, font="Helvetica-Oblique",
                                  textColor=colors.HexColor("#555555"))))
    story.append(Spacer(1, 5))

    meta = Table([["Name", "SRN", "Course"],
                  [STUDENT["name"], STUDENT["srn"], STUDENT["course"]]],
                 colWidths=[62 * mm, 40 * mm, 84 * mm])
    meta.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), HDR_RGB),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.4),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]))
    story.append(meta)
    story.append(Spacer(1, 5))

    rows = [[Paragraph("<b>%s</b>" % k, small), Paragraph(v, small)] for k, v in FLOW["meta"]]
    mt = Table(rows, colWidths=[34 * mm, 152 * mm])
    mt.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F4F7FC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]))
    story.append(mt)

    def numbered(items, style=body, indent=10):
        st = ParagraphStyle("n", parent=style, leftIndent=indent + 12, firstLineIndent=-12,
                            spaceAfter=1.2)
        for i, it in enumerate(items, start=1):
            story.append(Paragraph("<b>%d.</b>&nbsp;&nbsp;%s" % (i, it), st))

    def plain(items, indent=20):
        st = ParagraphStyle("p", parent=small, leftIndent=indent, spaceAfter=1.2)
        for it in items:
            story.append(Paragraph(it, st))

    story.append(Paragraph("Preconditions", sec))
    numbered(FLOW["preconditions"])
    story.append(Paragraph("Postconditions", sec))
    numbered(FLOW["postconditions"])
    story.append(Paragraph("Main Success Scenario", sec))
    numbered(FLOW["main"])
    story.append(Paragraph("Alternate Flows", sec))
    for alt in FLOW["alternates"]:
        story.append(Paragraph("<b>%s</b>" % alt["title"],
                               ParagraphStyle("at", parent=small, leftIndent=8,
                                              spaceBefore=3, spaceAfter=1.5)))
        plain(alt["steps"])
    story.append(Paragraph("Exception Flow", sec))
    plain([FLOW["exception"]], indent=10)

    doc.build(story)
    print("wrote docs/03_UseCase_Flow_UC-01.pdf")


if __name__ == "__main__":
    build_requirements_docx()
    build_requirements_xlsx()
    build_requirements_pdf()
    build_flow_docx()
    build_flow_pdf()
